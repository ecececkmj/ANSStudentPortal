"""
Creates an Access Code for everyone on the staff list who does not have one yet.

How to use:
  1. Open  AUTHORIZED STAFF.xlsx  and type each person's Full Name and Role (Teacher or SSLG).
     Leave the Access Code cell EMPTY.   Save and CLOSE Excel.
  2. Run this file (from any folder):   python make_codes.py
  3. The codes are written into the Excel file and printed here. Send each person THEIR OWN code privately.

Existing codes are never changed. (To give someone a new code, clear their cell and run this again.)
This script finds the Excel file by itself: it looks in the folder you are in, in the folder this
script is in, and in the folder above those (your project folder).
"""
import glob
import os
import re
import secrets
import sys
from decimal import Decimal, InvalidOperation

ALPHABET = 'ABCDEFGHJKMNPQRSTUVWXYZ23456789'      # no 0/O, 1/I/L so codes are easy to read out loud


def new_code():
    pick = lambda n: ''.join(secrets.choice(ALPHABET) for _ in range(n))
    return f'{pick(4)}-{pick(4)}'


def clean(value):
    if value is None:
        return ''
    if isinstance(value, (int, float)):
        try:
            return str(int(Decimal(str(value))))
        except (InvalidOperation, ValueError):
            return str(value).strip()
    text = str(value).strip()
    return text.split('.')[0] if re.fullmatch(r'\d+\.0+', text) else text


def find_staff_file():
    here = os.path.dirname(os.path.abspath(__file__))
    folders = []
    for f in (os.getcwd(), here, os.path.dirname(here), os.path.dirname(os.getcwd())):
        if f not in folders:
            folders.append(f)
    for folder in folders:
        found = []
        for path in glob.glob(os.path.join(folder, '*.xlsx')):
            name = os.path.basename(path).lower()
            if not name.startswith('~$') and ('staff' in name or 'authorized' in name):
                found.append(path)
        if len(found) == 1:
            return found[0]
        if len(found) > 1:
            sys.exit('More than one staff Excel found in ' + folder + ':\n  ' + '\n  '.join(found) +
                     '\nKeep only one file with "staff" or "authorized" in its name.')
    sys.exit('No staff Excel found. I looked in:\n  ' + '\n  '.join(folders) +
             '\nPut AUTHORIZED STAFF.xlsx in your project folder (next to app.py).')


def main():
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit('openpyxl is not installed. Run:  pip install openpyxl')

    path = find_staff_file()
    print('Using:', path)
    try:
        wb = load_workbook(path)
    except Exception as e:
        sys.exit(f'Could not open the file: {e}\nClose it in Excel and try again.')
    ws = wb.active

    used = {clean(ws.cell(r, 1).value).lower() for r in range(2, ws.max_row + 1)} - {''}
    made = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 2).value
        role = ws.cell(r, 3).value
        if not name or not role or clean(ws.cell(r, 1).value):
            continue
        code = new_code()
        while code.lower() in used:
            code = new_code()
        used.add(code.lower())
        ws.cell(r, 1).value = code
        made.append((str(name).strip(), str(role).strip(), code))

    if not made:
        print('Nobody needs a new code. (Fill in Full Name + Role for new people and leave their Access Code empty.)')
        return
    try:
        wb.save(path)
    except PermissionError:
        sys.exit('Could not save the file. Close it in Excel and run this again.')

    print(f'\nCreated {len(made)} access code(s) and saved them in the Excel file:\n')
    for name, role, code in made:
        print(f'  {name:<35} {role:<8} {code}')
    print('\nSend each person ONLY their own code, privately.')


if __name__ == '__main__':
    main()
